
import docx
import sys

def analyze_docx(filename):
    try:
        doc = docx.Document(filename)
        print(f"\n--- Analysis of {filename} ---")
        
        print("\n[Paragraphs]")
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip():
                print(f"{i}: {p.text[:100]}")
                if p.alignment:
                    print(f"   Alignment: {p.alignment}")
        
        print("\n[Tables]")
        for i, table in enumerate(doc.tables):
            print(f"Table {i+1} ({len(table.rows)} rows x {len(table.columns)} cols):")
            # Print column widths if available
            widths = []
            for cell in table.rows[0].cells:
                widths.append(cell.width)
            print(f"   Col widths: {widths}")
            
            for j, row in enumerate(table.rows):
                row_text = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                print(f"  Row {j}: {row_text}")

    except Exception as e:
        print(f"Error reading file {filename}: {e}")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        analyze_docx(sys.argv[1])
    else:
        print("Usage: python analyze_docx_v2.py <filename>")

