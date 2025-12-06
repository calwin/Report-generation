
import docx
import sys

try:
    doc = docx.Document('Report of the Standing Committee - R1.docx')
    print("--- Paragraphs ---")
    for p in doc.paragraphs:
        if p.text.strip():
            print(p.text[:100])
    
    print("\n--- Tables ---")
    for i, table in enumerate(doc.tables):
        print(f"Table {i+1}:")
        for row in table.rows:
            row_text = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            print(f"  {row_text}")

except Exception as e:
    print(f"Error reading file: {e}")

