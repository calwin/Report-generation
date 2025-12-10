#!/usr/bin/env python3
"""
Flask Backend for PQR Report Generation
Receives form data and generates Standing Committee Report
"""

from flask import Flask, request, jsonify, send_file, render_template, make_response
from flask_cors import CORS
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
import os
import io

app = Flask(__name__)
CORS(app)

class ReportGenerator:
    """Generate Standing Committee Report from form data"""

    def __init__(self, form_data):
        self.data = form_data
        self.doc = Document()

    def generate_text(self):
        """Generate plain text version for testing"""
        output = []

        # Title
        output.append("REPORT OF THE STANDING COMMITTEE FOR PRE-QUALIFICATION REQUIREMENT")
        output.append("")
        output.append("="*80)
        output.append("")

        # Row 1-4
        output.append(f"1. Name of the work: {self.data.get('work_name', 'N/A')}")
        output.append(f"2. Period of work: {self.data.get('work_period', 'N/A')}")
        output.append(f"3. Nature of work: {self.data.get('work_nature', 'AMC/BMC/Short Term')}")
        output.append(f"4. Estimate Value (Excl. GST): {self.data.get('estimate_value', 'EV')}")
        output.append("")

        # Row 5: Proposed PQR
        output.append("5. Proposed PQR:")
        pqr_content = self._build_pqr_text()
        output.append(pqr_content)
        output.append("")

        # Row 6-8
        output.append(f"6. Details of any similar work done earlier: {self.data.get('similar_work', '')}")
        output.append(f"7. If yes to Sl. No. 6, any change contemplated and brought out the reasons for change: {self.data.get('changes_contemplated', 'Financial requirements and Note modified in accordance with the PQR Standardisation circular.')}")
        output.append("")
        output.append("8. Recommendation of PQR Committee")
        output.append("   Reviewed and Recommended")
        output.append("")
        output.append("   Remarks Offered")
        output.append("")
        output.append("="*80)
        output.append("Deputy General Manager/BM\t\tDeputy General Manager/BM\t\tDeputy General Manager/BM\t\tDeputy General Manager/BM")

        return "\n".join(output)

    def _build_pqr_text(self):
        """Build PQR content as plain text"""
        lines = []

        # Section 1: Technical requirements
        lines.append("1. Technical requirements")

        tech_items = []
        if self.data.get('check_tech_exp') and self.data.get('tech_experience'):
            tech_items.append(self.data.get('tech_experience'))
        if self.data.get('check_nature_works') and self.data.get('nature_of_works'):
            tech_items.append(self.data.get('nature_of_works'))
        if self.data.get('check_location') and self.data.get('location_of_work'):
            tech_items.append(self.data.get('location_of_work'))
        if self.data.get('check_value_opt1') and self.data.get('value_option_1'):
            tech_items.append(self.data.get('value_option_1'))
        if self.data.get('check_value_opt2') and self.data.get('value_option_2'):
            tech_items.append(self.data.get('value_option_2'))
        if self.data.get('check_license') and self.data.get('license_requirement'):
            tech_items.append(self.data.get('license_requirement'))
        if self.data.get('check_gen_work') and self.data.get('general_work_experience'):
            tech_items.append(self.data.get('general_work_experience'))
        if self.data.get('check_general_license') and self.data.get('general_license'):
            tech_items.append(self.data.get('general_license'))

        for idx, item in enumerate(tech_items):
            letter = chr(97 + idx)
            lines.append(f"    {letter}. {item}")

        # Section 2: Financial Requirements
        lines.append("2. Financial Requirements:")
        fin_counter = 0
        if self.data.get('check_turnover') and self.data.get('turnover_requirement'):
            letter = chr(97 + fin_counter)
            lines.append(f"    {letter}. {self.data.get('turnover_requirement')}")
            fin_counter += 1
        if self.data.get('check_networth') and self.data.get('net_worth_requirement'):
            letter = chr(97 + fin_counter)
            lines.append(f"    {letter}. {self.data.get('net_worth_requirement')}")
            fin_counter += 1
        if self.data.get('check_fin_doc') and self.data.get('financial_documentation'):
            letter = chr(97 + fin_counter)
            lines.append(f"    {letter}. {self.data.get('financial_documentation')}")
            fin_counter += 1

        # Section 3: Note
        lines.append("3. Note:")
        lines.append("    a. Technical Requirements:")

        # Build clause references
        clause_refs = []
        for i in range(len(tech_items)):
            clause_refs.append(f'(1.{chr(97 + i)})')

        clause_string = ', '.join(clause_refs[:-1]) + ' & ' + clause_refs[-1] if len(clause_refs) > 1 else clause_refs[0] if clause_refs else ''

        if clause_string:
            lines.append(f"        The following documents or relevant documentary evidences so as to meet the above stipulated PQR Conditions {clause_string} are to be furnished along with the bid without fail. Otherwise, the offer is liable for rejection.")

        if self.data.get('req_work_order') and self.data.get('req_work_order_text'):
            lines.append(f"        {self.data.get('req_work_order_text')}")
        if self.data.get('req_completion_cert') and self.data.get('req_completion_cert_text'):
            lines.append(f"        {self.data.get('req_completion_cert_text')}")
        if self.data.get('req_license_copy') and self.data.get('req_license_copy_text'):
            lines.append(f"        {self.data.get('req_license_copy_text')}")
        if self.data.get('note_captive_exclusion') and self.data.get('note_captive_exclusion_text'):
            lines.append(f"        {self.data.get('note_captive_exclusion_text')}")

        lines.append("        The bidder shall be a Proprietary firm/Partnership firm/ a firm (Private, public, Govt) registered under Company's act.")

        if clause_refs:
            all_clauses = ', '.join([f'1.{chr(97+i)}' for i in range(len(tech_items))])
            lines.append(f"        The bidder should meet the criteria specified in PQR Clause {all_clauses}, 2 or Clause 1.b, {all_clauses[4:]}, 2.")

        lines.append("    b. Financial Requirements:")

        fin_clauses = []
        if self.data.get('check_turnover'):
            fin_clauses.append('2.a')
        if self.data.get('check_networth'):
            fin_clauses.append('2.b')

        fin_clause_string = ' & '.join(fin_clauses) if fin_clauses else 'financial clauses'
        lines.append(f"        Copies of standalone audited financial statement (Profit & Loss account and Balance Sheet) for three financial years immediately preceding the original scheduled tender opening date are to be furnished to meet the PQR mentioned in clause {fin_clause_string}. Otherwise the offer is liable for rejection.")

        if self.data.get('note_new_firm_turnover') and self.data.get('note_new_firm_turnover_text'):
            lines.append(f"        {self.data.get('note_new_firm_turnover_text')}")
        if self.data.get('note_turnover_definition') and self.data.get('note_turnover_definition_text'):
            lines.append(f"        {self.data.get('note_turnover_definition_text')}")
        if self.data.get('note_networth_companies') and self.data.get('note_networth_companies_text'):
            lines.append(f"        {self.data.get('note_networth_companies_text')}")
        if self.data.get('note_networth_proprietor') and self.data.get('note_networth_proprietor_text'):
            lines.append(f"        {self.data.get('note_networth_proprietor_text')}")
        if self.data.get('note_multi_agreement') and self.data.get('note_multi_agreement_text'):
            lines.append(f"        {self.data.get('note_multi_agreement_text')}")
        if self.data.get('note_exclude_gst') and self.data.get('note_exclude_gst_text'):
            lines.append(f"        {self.data.get('note_exclude_gst_text')}")
        if self.data.get('note_partnership') and self.data.get('note_partnership_text'):
            lines.append(f"        {self.data.get('note_partnership_text')}")

        return "\n".join(lines)

    def generate(self):
        """Generate the complete report and return as bytes"""

        # Title
        title = self.doc.add_paragraph()
        title_run = title.add_run('REPORT OF THE STANDING COMMITTEE FOR PRE-QUALIFICATION REQUIREMENT')
        title_run.bold = True
        title_run.font.size = Pt(12)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        self.doc.add_paragraph()  # Spacing

        # Create main table with 4 columns to match R1 format
        table = self.doc.add_table(rows=9, cols=4)
        table.style = 'Table Grid'

        # Set column widths
        from docx.shared import Inches
        table.columns[0].width = Inches(0.4)  # Number
        table.columns[1].width = Inches(2.0)  # Label
        table.columns[2].width = Inches(0.2)  # Separator (:)
        table.columns[3].width = Inches(4.5)  # Content

        # Row 1: Name of work
        self._add_table_row(table, 0, '1.', 'Name of the work',
                           self.data.get('work_name', 'N/A'))

        # Row 2: Period of work
        self._add_table_row(table, 1, '2', 'Period of work',
                           self.data.get('work_period', 'N/A'))

        # Row 3: Nature of work
        nature = self.data.get('work_nature', 'AMC/BMC/Short Term')
        self._add_table_row(table, 2, '3.', 'Nature of work', nature)

        # Row 4: Estimate Value
        estimate_value = self.data.get('estimate_value', 'EV')
        self._add_table_row(table, 3, '4.', 'Estimate Value (Excl. GST)', estimate_value)

        # Row 5: Proposed PQR
        row5_cells = table.rows[4].cells
        row5_cells[0].text = '5'
        row5_cells[1].text = 'Proposed PQR'
        row5_cells[2].text = ':'
        
        # Build PQR content with proper formatting
        pqr_paragraph = row5_cells[3].paragraphs[0]
        pqr_paragraph.text = ''
        
        # Build and add formatted PQR content
        self._add_formatted_pqr_content(pqr_paragraph)

        # Row 6: Details of similar work
        similar_work = self.data.get('similar_work', '')
        self._add_table_row(table, 5, '6', 'Details of any similar work done earlier', similar_work)

        # Row 7: Changes contemplated
        changes = self.data.get('changes_contemplated',
                               'Financial requirements and Note modified in accordance with the PQR Standardisation circular.')
        self._add_table_row(table, 6, '7',
                          'If yes to Sl. No. 6, any change contemplated and brought out the reasons for change',
                          changes)

        # Row 8: Signatures
        row7 = table.rows[7]
        # Merge all cells for signature row to avoid narrow column wrapping
        row7.cells[0].merge(row7.cells[3])
        # Align to right as per example screenshot
        row7.cells[0].text = 'Deputy General Manager/BM'
        row7.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Row 9: Recommendations
        row8 = table.rows[8]
        row8.cells[0].text = '8'
        
        # Recommendation text with checkboxes
        # Merge remaining columns for content
        row8.cells[1].merge(row8.cells[3])
        rec_para = row8.cells[1].paragraphs[0]
        rec_para.add_run('Recommendation of PQR Committee\n')
        rec_para.add_run('\n')
        rec_para.add_run(u'\u2610 Reviewed and Recommended\n')
        rec_para.add_run('\n')
        rec_para.add_run(u'\u2610 Remarks Offered')

        # Add Member signature lines at bottom using a table for proper alignment
        self.doc.add_paragraph() # Spacing
        
        # Create a 3-column table for signatures (no borders by default if style not set to 'Table Grid')
        member_table = self.doc.add_table(rows=1, cols=3)
        member_table.autofit = True
        
        # Left Member
        cell_left = member_table.rows[0].cells[0]
        p_left = cell_left.paragraphs[0]
        run = p_left.add_run('Member')
        run.bold = True
        p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        # Center Member
        cell_center = member_table.rows[0].cells[1]
        p_center = cell_center.paragraphs[0]
        run = p_center.add_run('Member')
        run.bold = True
        p_center.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Right Member
        cell_right = member_table.rows[0].cells[2]
        p_right = cell_right.paragraphs[0]
        run = p_right.add_run('Member')
        run.bold = True
        p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Save to bytes
        doc_io = io.BytesIO()
        self.doc.save(doc_io)
        doc_io.seek(0)
        return doc_io

    def _add_table_row(self, table, row_idx, num, label, value):
        """Helper to add a row to the table"""
        row = table.rows[row_idx]
        row.cells[0].text = num
        row.cells[1].text = label
        row.cells[2].text = ':'
        row.cells[3].text = str(value)

    def _add_formatted_pqr_content(self, paragraph):
        """Add formatted PQR content with proper paragraph numbering like the example"""

        # Section 1: Technical requirements (bold header with number)
        run = paragraph.add_run('1. ')
        run = paragraph.add_run('Technical requirements')
        run.bold = True
        paragraph.add_run('\n')

        # Track which technical requirements are checked and add with letter numbering
        tech_items = []

        if self.data.get('check_tech_exp') and self.data.get('tech_experience'):
            tech_items.append(self.data.get('tech_experience'))

        if self.data.get('check_nature_works') and self.data.get('nature_of_works'):
            tech_items.append(self.data.get('nature_of_works'))

        if self.data.get('check_location') and self.data.get('location_of_work'):
            tech_items.append(self.data.get('location_of_work'))

        if self.data.get('check_value_opt1') and self.data.get('value_option_1'):
            tech_items.append(self.data.get('value_option_1'))

        if self.data.get('check_value_opt2') and self.data.get('value_option_2'):
            tech_items.append(self.data.get('value_option_2'))

        if self.data.get('check_license') and self.data.get('license_requirement'):
            tech_items.append(self.data.get('license_requirement'))

        if self.data.get('check_gen_work') and self.data.get('general_work_experience'):
            tech_items.append(self.data.get('general_work_experience'))

        if self.data.get('check_general_license') and self.data.get('general_license'):
            tech_items.append(self.data.get('general_license'))

        # Add technical items with letter prefixes
        for idx, item in enumerate(tech_items):
            letter = chr(97 + idx)  # a, b, c, d, e...
            # Handle multiline items by splitting and indenting subsequent lines
            lines = str(item).split('\n')
            paragraph.add_run(f'    {letter}. {lines[0]}\n')
            for line in lines[1:]:
                if line.strip():
                    paragraph.add_run(f'       {line.strip()}\n')

        # Section 2: Financial Requirements (bold header with number)
        run = paragraph.add_run('2. ')
        run = paragraph.add_run('Financial Requirements:')
        run.bold = True
        paragraph.add_run('\n')

        # Add financial requirements with letter numbering
        fin_counter = 0
        if self.data.get('check_turnover') and self.data.get('turnover_requirement'):
            letter = chr(97 + fin_counter)
            paragraph.add_run(f'    {letter}. {self.data.get("turnover_requirement")}\n')
            fin_counter += 1

        if self.data.get('check_networth') and self.data.get('net_worth_requirement'):
            letter = chr(97 + fin_counter)
            paragraph.add_run(f'    {letter}. {self.data.get("net_worth_requirement")}\n')
            fin_counter += 1

        if self.data.get('check_fin_doc') and self.data.get('financial_documentation'):
            letter = chr(97 + fin_counter)
            paragraph.add_run(f'    {letter}. {self.data.get("financial_documentation")}\n')
            fin_counter += 1

        # Section 3: Note (bold header with number)
        run = paragraph.add_run('3. ')
        run = paragraph.add_run('Note:')
        run.bold = True
        paragraph.add_run('\n')

        # Note subsection a: Technical Requirements (bold)
        run = paragraph.add_run('    a. ')
        run = paragraph.add_run('Technical Requirements:')
        run.bold = True
        paragraph.add_run('\n')

        # Build the clause references dynamically based on what's checked
        clause_refs = []
        tech_counter = 0
        for i in range(len(tech_items)):
            tech_counter += 1
            clause_refs.append(f'(1.{chr(96 + tech_counter)})')  # (1.a), (1.b), etc.

        clause_string = ', '.join(clause_refs[:-1]) + ' & ' + clause_refs[-1] if len(clause_refs) > 1 else clause_refs[0] if clause_refs else ''

        if clause_string:
            paragraph.add_run(f'        The following documents or relevant documentary evidences so as to meet the above stipulated PQR Conditions {clause_string} are to be furnished along with the bid without fail. Otherwise, the offer is liable for rejection.\n')
        else:
            paragraph.add_run('        The following documents or relevant documentary evidences so as to meet the above stipulated PQR Conditions are to be furnished along with the bid without fail. Otherwise, the offer is liable for rejection.\n')

        # Documentation requirements
        if self.data.get('req_work_order') and self.data.get('req_work_order_text'):
            paragraph.add_run(f'        {self.data.get("req_work_order_text")}\n')

        if self.data.get('req_completion_cert') and self.data.get('req_completion_cert_text'):
            paragraph.add_run(f'        {self.data.get("req_completion_cert_text")}\n')

        if self.data.get('req_license_copy') and self.data.get('req_license_copy_text'):
            paragraph.add_run(f'        {self.data.get("req_license_copy_text")}\n')

        if self.data.get('note_captive_exclusion') and self.data.get('note_captive_exclusion_text'):
            paragraph.add_run(f'        {self.data.get("note_captive_exclusion_text")}\n')

        # Standard notes
        paragraph.add_run('        The bidder shall be a Proprietary firm/Partnership firm/ a firm (Private, public, Govt) registered under Company\'s act.\n')

        # Build criteria specification dynamically
        if clause_refs:
            # Build two options for criteria
            all_clauses = ', '.join([f'1.{chr(97+i)}' for i in range(len(tech_items))])
            paragraph.add_run(f'        The bidder should meet the criteria specified in PQR Clause {all_clauses}, 2 or Clause 1.b, {all_clauses[4:]}, 2.\n')
        else:
            paragraph.add_run('        The bidder should meet the criteria specified in the applicable PQR clauses.\n')

        # Note subsection b: Financial Requirements (bold)
        run = paragraph.add_run('    b. ')
        run = paragraph.add_run('Financial Requirements:')
        run.bold = True
        paragraph.add_run('\n')

        # Build financial clause references
        fin_clauses = []
        if self.data.get('check_turnover'):
            fin_clauses.append('2.a')
        if self.data.get('check_networth'):
            fin_clauses.append('2.b')

        fin_clause_string = ' & '.join(fin_clauses) if fin_clauses else 'financial clauses'

        paragraph.add_run(f'        Copies of standalone audited financial statement (Profit & Loss account and Balance Sheet) for three financial years immediately preceding the original scheduled tender opening date are to be furnished to meet the PQR mentioned in clause {fin_clause_string}. Otherwise the offer is liable for rejection.\n')

        # Additional notes if checked
        if self.data.get('note_new_firm_turnover') and self.data.get('note_new_firm_turnover_text'):
            paragraph.add_run(f'        {self.data.get("note_new_firm_turnover_text")}\n')

        if self.data.get('note_turnover_definition') and self.data.get('note_turnover_definition_text'):
            paragraph.add_run(f'        {self.data.get("note_turnover_definition_text")}\n')

        if self.data.get('note_networth_companies') and self.data.get('note_networth_companies_text'):
            paragraph.add_run(f'        {self.data.get("note_networth_companies_text")}\n')

        if self.data.get('note_networth_proprietor') and self.data.get('note_networth_proprietor_text'):
            paragraph.add_run(f'        {self.data.get("note_networth_proprietor_text")}\n')

        if self.data.get('note_multi_agreement') and self.data.get('note_multi_agreement_text'):
            paragraph.add_run(f'        {self.data.get("note_multi_agreement_text")}\n')

        if self.data.get('note_exclude_gst') and self.data.get('note_exclude_gst_text'):
            paragraph.add_run(f'        {self.data.get("note_exclude_gst_text")}\n')

        if self.data.get('note_partnership') and self.data.get('note_partnership_text'):
            paragraph.add_run(f'        {self.data.get("note_partnership_text")}\n')


@app.route('/')
def index():
    """Serve the main form page"""
    return render_template('index.html')


@app.route('/api/generate-report', methods=['POST'])
def generate_report():
    """API endpoint to generate report from form data"""
    try:
        # Get form data from request
        form_data = request.json

        if not form_data:
            return jsonify({'error': 'No data provided'}), 400

        # Generate report as DOCX
        generator = ReportGenerator(form_data)
        doc_bytes = generator.generate()

        # Create response with DOCX file
        response = make_response(doc_bytes)
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        response.headers['Content-Disposition'] = 'attachment; filename=Standing_Committee_Report.docx'

        return response

    except Exception as e:
        import traceback
        print(f"ERROR: {str(e)}")
        print(traceback.format_exc())
        return jsonify({'error': str(e)}), 500


@app.route('/api/health', methods=['GET'])
def health_check():
    """Health check endpoint"""
    return jsonify({'status': 'healthy', 'timestamp': datetime.now().isoformat()})


if __name__ == '__main__':
    # Create templates directory if it doesn't exist
    os.makedirs('templates', exist_ok=True)

    print("=" * 80)
    print("PQR REPORT GENERATOR - Starting Flask Server")
    print("=" * 80)
    
    # Use PORT from environment (Railway) or default to 8001
    port = int(os.environ.get('PORT', 8001))
    print(f"\nServer will be available at: http://localhost:{port}")
    print("\nPress CTRL+C to stop the server")
    print("=" * 80)

    app.run(debug=True, host='0.0.0.0', port=port)
